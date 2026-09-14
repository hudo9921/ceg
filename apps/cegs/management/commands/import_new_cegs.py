import os
import re
from datetime import datetime, timedelta
from decimal import Decimal
import openpyxl
from django.core.management.base import BaseCommand
from django.utils import timezone
from django.utils.text import slugify
from django.contrib.auth.models import User
from django.db import transaction

from apps.groups.models import KpopGroup, Era
from apps.cegs.models import (
    CEG,
    CEGItemDefinition,
    CEGSet,
    ItemSlot,
    TipoItem,
    ClaimAttemptLog,
    ItemWaitingList,
    Caixa,
    CaixaItemRate,
    ItemIndividual,
    PacoteNacional,
)
from apps.participants.models import Participant, Claim, ParticipantNotification


class Command(BaseCommand):
    help = (
        'Limpa a massa operacional anterior e importa as novas planilhas de CEGs '
        '(Made My Night e Pureflow) com todas as vagas limpas e disponíveis.'
    )

    def add_arguments(self, parser):
        parser.add_argument(
            '--mmn-file',
            type=str,
            default='planilhas/CEGs_Lessera_Made_My_Night_Reestruturada.xlsx',
            help='Caminho para o arquivo da planilha Made My Night (.xlsx)'
        )
        parser.add_argument(
            '--pf-file',
            type=str,
            default='planilhas/CEGs_Lessera_Pureflow_Reestruturada.xlsx',
            help='Caminho para o arquivo da planilha Pureflow (.xlsx)'
        )
        parser.add_argument(
            '--clear-participants',
            action='store_true',
            default=False,
            help='Se fornecido, também apaga todos os participantes cadastrados.'
        )

    def parse_prazo(self, raw_prazo, now, tz):
        """Converte strings no formato 'DD/MM' para datetime com ano 2026."""
        if not raw_prazo:
            return None
        raw_str = str(raw_prazo).strip()
        m = re.match(r'(\d{1,2})/(\d{1,2})', raw_str)
        if m:
            day, month = int(m.group(1)), int(m.group(2))
            return datetime(2026, month, day, 23, 59, 59, tzinfo=tz)
        return None

    def handle(self, *args, **options):
        mmn_path = options['mmn_file']
        pf_path = options['pf_file']
        clear_participants = options['clear_participants']

        if not os.path.exists(mmn_path):
            self.stderr.write(self.style.ERROR(f"Arquivo não encontrado: {mmn_path}"))
            return
        if not os.path.exists(pf_path):
            self.stderr.write(self.style.ERROR(f"Arquivo não encontrado: {pf_path}"))
            return

        self.stdout.write(self.style.MIGRATE_HEADING("=== IMPORTAÇÃO DE NOVAS PLANILHAS DE CEGS (LIMPAS) ==="))

        with transaction.atomic():
            # 1. Limpeza da massa de dados operacional antiga
            self.stdout.write("1. Limpando massa operacional antiga...")
            ItemWaitingList.objects.all().delete()
            ClaimAttemptLog.objects.all().delete()
            Claim.objects.all().delete()
            ItemSlot.objects.all().delete()
            CEGSet.objects.all().delete()
            CEGItemDefinition.objects.all().delete()
            CaixaItemRate.objects.all().delete()
            ItemIndividual.objects.all().delete()
            PacoteNacional.objects.all().delete()
            CEG.objects.all().delete()
            Caixa.objects.all().delete()
            ParticipantNotification.objects.all().delete()

            if clear_participants:
                Participant.objects.all().delete()
                self.stdout.write("   Participantes removidos conforme solicitado.")
            else:
                self.stdout.write(f"   Participantes preservados: {Participant.objects.count()} cadastrados.")

            self.stdout.write(self.style.SUCCESS("   Massa anterior de CEGs/Slots/Claims excluída com sucesso."))

            # 2. Superusuário admin
            admin_user, created_admin = User.objects.get_or_create(
                username='admin',
                defaults={
                    'email': 'admin@example.com',
                    'is_staff': True,
                    'is_superuser': True,
                }
            )
            if created_admin:
                admin_user.set_password('admin123')
                admin_user.save()
                self.stdout.write(self.style.SUCCESS("   Superusuário criado: login 'admin' | senha 'admin123'"))

            # 3. Tipos de Item Padrão
            tipo_photocard, _ = TipoItem.objects.get_or_create(
                nome='Photocard',
                defaults={'descricao': 'Cards promocionais, POB ou de álbum'}
            )
            tipo_album, _ = TipoItem.objects.get_or_create(
                nome='Álbum',
                defaults={'descricao': 'Álbuns completos ou compactos'}
            )
            tipo_cd, _ = TipoItem.objects.get_or_create(
                nome='CD',
                defaults={'descricao': 'CDs e mídias'}
            )

            # 4. Grupo K-Pop LE SSERAFIM
            group, _ = KpopGroup.objects.get_or_create(
                name='LE SSERAFIM',
                defaults={
                    'slug': 'le-sserafim',
                    'image_url': 'https://images.unsplash.com/photo-1514525253161-7a46d19cd819?w=600&auto=format&fit=crop',
                    'description': 'Girl group sul-coreano de 5 integrantes da Source Music / HYBE.'
                }
            )

            # 5. Eras: Made My Night e Pureflow
            era_mmn, created_mmn = Era.objects.get_or_create(
                group=group,
                name='Made My Night',
                defaults={
                    'slug': 'made-my-night',
                    'release_date': datetime(2026, 9, 1).date(),
                    'description': 'Era Made My Night — compras em grupo de POBs, Broadcast, Lucky Draws e álbuns.'
                }
            )
            era_pf, created_pf = Era.objects.get_or_create(
                group=group,
                name='Pureflow',
                defaults={
                    'slug': 'pureflow',
                    'release_date': datetime(2026, 5, 30).date(),
                    'description': 'Álbum e Especiais Pureflow de LE SSERAFIM.'
                }
            )
            self.stdout.write(
                f"2. Eras verificadas/criadas: '{era_mmn.name}' ({'nova' if created_mmn else 'existente'}) "
                f"e '{era_pf.name}' ({'nova' if created_pf else 'existente'})."
            )

            tz = timezone.get_current_timezone()
            now = timezone.now()

            # 6. Função de Importação de Planilha
            def import_workbook(path, era_obj):
                self.stdout.write(f"\n-> Processando: {path} (Era: {era_obj.name})")
                wb = openpyxl.load_workbook(path, data_only=True)

                # Leitura de CEGs_Cadastro
                ws_ceg = wb['CEGs_Cadastro']
                ceg_registry = {}  # ceg_id -> dict com metadados e objeto CEG

                for r in range(2, ws_ceg.max_row + 1):
                    ceg_id_raw = ws_ceg.cell(r, 1).value
                    if not ceg_id_raw:
                        continue
                    ceg_id = str(ceg_id_raw).strip()
                    ceg_name = str(ws_ceg.cell(r, 3).value or '').strip()

                    raw_valor_item = ws_ceg.cell(r, 4).value
                    valor_item = Decimal(str(raw_valor_item or '0.00'))

                    raw_prazo_item = ws_ceg.cell(r, 5).value
                    prazo_item_dt = self.parse_prazo(raw_prazo_item, now, tz) or (now + timedelta(days=7))

                    raw_valor_inter = ws_ceg.cell(r, 6).value
                    valor_inter = Decimal(str(raw_valor_inter)) if raw_valor_inter not in (None, '') else None

                    raw_prazo_inter = ws_ceg.cell(r, 7).value
                    prazo_inter_dt = self.parse_prazo(raw_prazo_inter, now, tz)

                    raw_valor_taxa = ws_ceg.cell(r, 8).value
                    valor_taxa = Decimal(str(raw_valor_taxa)) if raw_valor_taxa not in (None, '') else None

                    raw_prazo_taxa = ws_ceg.cell(r, 9).value
                    prazo_taxa_dt = self.parse_prazo(raw_prazo_taxa, now, tz)

                    slug = f"ceg-{slugify(ceg_id)}-{slugify(ceg_name)}"

                    ceg_obj, _ = CEG.objects.get_or_create(
                        slug=slug,
                        defaults={
                            'era': era_obj,
                            'title': f"LE SSERAFIM — {ceg_name} ({ceg_id})",
                            'description': (
                                f"Compra em Grupo oficial para a era {era_obj.name} ({ceg_name}).\n"
                                f"• Prazo de pagamento do slot: {prazo_item_dt.strftime('%d/%m/%Y às %H:%M')}.\n"
                                "• Frete internacional e taxa aduaneira serão rateados na chegada da caixa."
                            ),
                            'pix_key': 'lessera.ceg@kpopbrasil.com.br',
                            'pix_instructions': 'Chave Pix (E-mail). Por favor anexar comprovante no WhatsApp.',
                            'status': CEG.Status.OPEN,
                            'opens_at': now - timedelta(days=2),
                            'closes_at': None,
                            'prazo_pagamento_item': prazo_item_dt,
                            'frete_inter': valor_inter,
                            'prazo_pagamento_frete_inter': prazo_inter_dt,
                            'taxa_aduaneira': valor_taxa,
                            'prazo_pagamento_taxa_aduaneira': prazo_taxa_dt,
                        }
                    )

                    ceg_registry[ceg_id] = {
                        'ceg_obj': ceg_obj,
                        'ceg_name': ceg_name,
                        'valor_item': valor_item,
                        'prazo_item_dt': prazo_item_dt,
                    }

                # Leitura de Vagas_Operacional
                ws_vagas = wb['Vagas_Operacional']

                # Passo A: Mapear membros e sets por CEG preservando a ordem
                ceg_members_order = {}  # ceg_id -> list of members
                ceg_sets_rows = {}      # ceg_id -> list of (set_num, status_set, member)

                for r in range(2, ws_vagas.max_row + 1):
                    ceg_id_raw = ws_vagas.cell(r, 1).value
                    if not ceg_id_raw:
                        continue
                    ceg_id = str(ceg_id_raw).strip()
                    if ceg_id not in ceg_registry:
                        continue

                    set_raw = str(ws_vagas.cell(r, 3).value or 'SET 01').strip().upper()
                    m_set = re.search(r'\d+', set_raw)
                    set_num = int(m_set.group()) if m_set else 1

                    status_set = str(ws_vagas.cell(r, 4).value or '').strip()
                    member = str(ws_vagas.cell(r, 5).value or '').strip()

                    if ceg_id not in ceg_members_order:
                        ceg_members_order[ceg_id] = []
                    if member and member not in ceg_members_order[ceg_id]:
                        ceg_members_order[ceg_id].append(member)

                    if ceg_id not in ceg_sets_rows:
                        ceg_sets_rows[ceg_id] = []
                    ceg_sets_rows[ceg_id].append((set_num, status_set, member))

                # Passo B: Criar CEGItemDefinition para cada membro da CEG
                total_slots_imported = 0
                total_sets_imported = 0

                for ceg_id, reg in ceg_registry.items():
                    ceg_obj = reg['ceg_obj']
                    ceg_name = reg['ceg_name']
                    valor_item = reg['valor_item']

                    is_album = 'compact' in ceg_name.lower() or 'álbum' in ceg_name.lower() or 'album' in ceg_name.lower()
                    chosen_tipo = tipo_album if is_album else tipo_photocard
                    chosen_item_type = CEGItemDefinition.ItemType.ALBUM if is_album else (
                        CEGItemDefinition.ItemType.POB if 'pob' in ceg_name.lower() else (
                            CEGItemDefinition.ItemType.OTHER if 'lucky draw' in ceg_name.lower() or 'ld' in ceg_name.lower() else CEGItemDefinition.ItemType.PHOTOCARD
                        )
                    )

                    members_list = ceg_members_order.get(ceg_id, [])
                    if not members_list:
                        members_list = ['Chaewon', 'Sakura', 'Yunjin', 'Kazuha', 'Eunchae']

                    item_defs_map = {}
                    for idx, member in enumerate(members_list, 1):
                        # Nome amigável do item
                        if is_album:
                            def_name = f"Compact ver. {member}"
                        elif member.lower().startswith('sakura ') or member.lower() == 'ot5':
                            def_name = f"Photocard {member}"
                        else:
                            def_name = f"Photocard {member}"

                        item_def, _ = CEGItemDefinition.objects.get_or_create(
                            ceg=ceg_obj,
                            member_name=member,
                            defaults={
                                'name': def_name,
                                'item_type': chosen_item_type,
                                'tipo_item': chosen_tipo,
                                'default_price': valor_item,
                                'order_index': idx,
                            }
                        )
                        item_defs_map[member] = item_def

                    # Passo C: Criar CEGSets e ItemSlots
                    sets_cache = {}
                    for set_num, status_set, member in ceg_sets_rows.get(ceg_id, []):
                        if set_num not in sets_cache:
                            cset, created_set = CEGSet.objects.get_or_create(
                                ceg=ceg_obj,
                                set_number=set_num,
                                defaults={
                                    'is_active': True,
                                    'notes': status_set,
                                }
                            )
                            if created_set:
                                total_sets_imported += 1
                            sets_cache[set_num] = cset
                        cset = sets_cache[set_num]

                        item_def = item_defs_map.get(member)
                        if not item_def:
                            continue

                        # Cria slot disponível (limpo, sem claims)
                        slot, created_slot = ItemSlot.objects.get_or_create(
                            set=cset,
                            item_definition=item_def,
                            defaults={
                                'price': valor_item,
                                'status': ItemSlot.Status.AVAILABLE,
                                'claimed_by': None,
                                'is_item_paid': False,
                                'is_frete_inter_paid': False,
                                'is_taxa_aduaneira_paid': False,
                            }
                        )
                        if created_slot:
                            total_slots_imported += 1

                self.stdout.write(self.style.SUCCESS(
                    f"   Concluído: {len(ceg_registry)} CEGs, {total_sets_imported} Sets e "
                    f"{total_slots_imported} Slots gerados (100% disponíveis)."
                ))
                return len(ceg_registry), total_sets_imported, total_slots_imported

            # Executa a importação das duas planilhas
            c1, s1, v1 = import_workbook(mmn_path, era_mmn)
            c2, s2, v2 = import_workbook(pf_path, era_pf)

        self.stdout.write("\n" + "=" * 60)
        self.stdout.write(self.style.SUCCESS(
            f"IMPORTAÇÃO CONCLUÍDA COM SUCESSO!\n"
            f"• Total de CEGs Cadastradas: {c1 + c2} ({c1} MMN + {c2} Pureflow)\n"
            f"• Total de Sets: {s1 + s2}\n"
            f"• Total de Vagas/Slots Disponíveis: {v1 + v2}\n"
            f"• Todas as vagas estão livres para claims ou alocação em massa!"
        ))
        self.stdout.write("=" * 60)
